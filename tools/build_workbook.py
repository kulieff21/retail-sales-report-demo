"""Excel workbook -> reports/retail-sales-report.xlsx. Every number is read from results/*.json.

Usage: uv run python tools/build_workbook.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from retail_report import report_style as S
from retail_report.report_data import (CITATION, DEMO_NOTE, LIMITATIONS, REPORTS, SOURCE_LINE, D, answers, facts,
                                       ledger_rows, ledger_total, load_all, year_pairs)

GBP = '£#,##0;[Red]-£#,##0'
GBP2 = '£#,##0.00;[Red]-£#,##0.00'
PCT = "0.0%"
PCT2 = "0.00%"
INT = "#,##0"
FONT = "Arial"

INK = S.INK.lstrip("#").upper()
f_title = Font(name=FONT, size=16, bold=True, color=INK)
f_sub = Font(name=FONT, size=10, color=S.MUTED.lstrip("#").upper())
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_body = Font(name=FONT, size=10, color=INK)
f_bold = Font(name=FONT, size=10, bold=True, color=INK)
f_sec = Font(name=FONT, size=12, bold=True, color=INK)
fill_head = PatternFill("solid", fgColor=INK)
fill_total = PatternFill("solid", fgColor="EEF1F4")
fill_bg = PatternFill("solid", fgColor="FFFEFB")
rule = Side(style="thin", color=S.RULE.lstrip("#").upper())
b_bottom = Border(bottom=rule)
b_top = Border(top=Side(style="thin", color=INK))


def num(x) -> float:
    return float(D(x))


class Sheet:
    """Small writer that tracks the current row and applies the shared styling."""

    def __init__(self, wb: Workbook, title: str, widths: list[float], tab: str | None = None):
        self.ws = wb.create_sheet(title)
        self.ws.sheet_view.showGridLines = False
        if tab:
            self.ws.sheet_properties.tabColor = tab
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.r = 1

    def title(self, text: str, sub: str | None = None):
        self.ws.cell(self.r, 1, text).font = f_title
        self.ws.row_dimensions[self.r].height = 24
        self.r += 1
        if sub:
            self.ws.cell(self.r, 1, sub).font = f_sub
            self.r += 1
        self.r += 1

    def section(self, text: str, note: str | None = None):
        self.ws.cell(self.r, 1, text).font = f_sec
        self.r += 1
        if note:
            self.ws.cell(self.r, 1, note).font = f_sub
            self.r += 1

    def header(self, cols: list[str]) -> int:
        for j, c in enumerate(cols, 1):
            cell = self.ws.cell(self.r, j, c)
            cell.font, cell.fill = f_head, fill_head
            cell.alignment = Alignment(horizontal="left" if j == 1 else "right", vertical="center", wrap_text=True)
        self.ws.row_dimensions[self.r].height = 30
        self.r += 1
        return self.r - 1

    def row(self, values: list, fmts: list[str | None] | None = None, bold: bool = False, total: bool = False) -> int:
        for j, v in enumerate(values, 1):
            cell = self.ws.cell(self.r, j, v)
            cell.font = f_bold if (bold or total) else f_body
            if fmts and j - 1 < len(fmts) and fmts[j - 1] and isinstance(v, (int, float)):
                cell.number_format = fmts[j - 1]
            cell.alignment = Alignment(horizontal="left" if isinstance(v, str) else "right", vertical="center")
            cell.border = b_top if total else b_bottom
            if total:
                cell.fill = fill_total
        self.r += 1
        return self.r - 1

    def text(self, text: str, font=None, wrap_cols: int = 1, height: float | None = None):
        c = self.ws.cell(self.r, 1, text)
        c.font = font or f_body
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if wrap_cols > 1:
            self.ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=wrap_cols)
        if height:
            self.ws.row_dimensions[self.r].height = height
        self.r += 1

    def gap(self, n: int = 1):
        self.r += n

    def freeze(self, row: int, col: int = 1):
        self.ws.freeze_panes = self.ws.cell(row, col)


def height_for(text: str, chars_per_line: int) -> float:
    lines = max(1, -(-len(text) // chars_per_line))
    return 14.5 * lines + 2


# ---------------------------------------------------------------- sheets

def sheet_readme(wb, R, F):
    s = Sheet(wb, "README", [118], tab=INK)
    s.title("Retail sales report (demo)", "Where did revenue growth come from, how dependent is it on a few customers, and what do cancellations cost?")
    W = 118
    def para(t, font=None):
        s.text(t, font, height=height_for(t, W - 6))
    s.section("Source")
    para("Online Retail II, UCI Machine Learning Repository. Licensed under CC BY 4.0; the raw file is not redistributed with this workbook.")
    para("Citation: " + CITATION)
    para(f"{DEMO_NOTE}: a portfolio piece built to show method, not a study of any real client. One retailer, not a market.", f_bold)
    s.gap()
    s.section("Definitions")
    for t in (
        "Gross product sales (GPS): sum of quantity x price over product sale lines.",
        "Cancelled product value (CPV): sum of the absolute value of product cancellation lines.",
        "Net product revenue (NPR): GPS minus CPV, each line dated on its own invoice date.",
        "Year A: Dec 2009 - Nov 2010. Year B: Dec 2010 - Nov 2011. December 2011 (1-9 Dec) is partial and excluded from every comparison.",
        "Order: a distinct numeric invoice with at least one product sale line.",
        "Customer status: retained = bought in both years; new = bought in Year B only; lost = bought in Year A only; "
        "cancel-only = identified customers with product cancellations but no sale in either year; no customer ID = lines with a missing customer ID.",
        "Repeat rate: share of a cohort placing a second order on a later date within 90 days of the first order (cohorts 2010-03 to 2011-08 only).",
        "No cost data exists, so nothing in this workbook is about margin or profit.",
    ):
        para(t)
    s.gap()
    s.section("Sheets")
    for t in ("Summary: headline numbers and the story in four sentences.",
              "Monthly: NPR, GPS and CPV by month, with a Year A vs Year B line chart.",
              "Growth: NPR change by country group, customer status and product, plus the customer bridge.",
              "Customers: concentration, repeat purchase by cohort and the retention matrix.",
              "Cancellations: monthly CPV with and without the 10 largest lines, those lines, product cancellation rates and traceability.",
              "Ledger: every raw row has exactly one disposition; rows and GBP reconcile to the raw total.",
              "Checks: independent re-derivation of the numbers."):
        para(t)
    s.gap()
    s.section("Limitations")
    for t in LIMITATIONS:
        para("- " + t)
    s.gap()
    s.section("Reproduce")
    for t in ("1. Get online_retail_II.xlsx from the UCI page above (SHA-256 is checked by tools/export_raw_csv.py).",
              "2. uv run python tools/export_raw_csv.py, then uv run python -m retail_report.build_results (writes results/*.json).",
              "3. uv run python tools/build_reports.py (figures, this workbook, the summary slide and the notebook).",
              "Every number in this workbook is read from results/*.json by tools/build_workbook.py; none is typed by hand."):
        para(t)
    s.gap()
    para(SOURCE_LINE, f_sub)


def sheet_summary(wb, R, F):
    q1, q3 = R["q1"], R["q3"]
    br = q1["bridge"]
    s = Sheet(wb, "Summary", [46, 17, 17, 17, 14], tab=S.GROWTH.lstrip("#"))
    s.title("Summary", "Year A = Dec 2009-Nov 2010, Year B = Dec 2010-Nov 2011. Amounts in GBP.")
    s.section("Headline numbers")
    s.header(["Metric", "Year A", "Year B", "Change", "Change %"])
    ha, hb = R["rec"]["section3"]["year_a"], R["rec"]["section3"]["year_b"]
    def money_row(label, a, b):
        a, b = num(a), num(b)
        return s.row([label, a, b, b - a, (b - a) / a if a else None], [None, GBP, GBP, GBP, PCT])
    money_row("Net product revenue (NPR)", ha["npr_gbp"], hb["npr_gbp"])
    money_row("Gross product sales (GPS)", ha["gps_gbp"], hb["gps_gbp"])
    money_row("Cancelled product value (CPV)", ha["cpv_gbp"], hb["cpv_gbp"])
    s.row(["CPV / GPS", F["cpv_rate_a"], F["cpv_rate_b"], F["cpv_rate_b"] - F["cpv_rate_a"], None], [None, PCT, PCT, PCT2])
    s.row(["CPV / GPS without the 10 largest lines", F["cpv_rate_a_wo"], F["cpv_rate_b_wo"], F["cpv_rate_b_wo"] - F["cpv_rate_a_wo"], None], [None, PCT, PCT, PCT2])
    money_row("Identified-customer NPR", br["year_a"]["identified_npr_gbp"], br["year_b"]["identified_npr_gbp"])
    money_row("No-customer-ID NPR", br["year_a"]["no_id_npr_gbp"], br["year_b"]["no_id_npr_gbp"])
    for lab, k in (("Orders", "orders"), ("Identified customers", "customers")):
        a, b = ha[k if k == "orders" else "identified_customers"], hb[k if k == "orders" else "identified_customers"]
        s.row([lab, a, b, b - a, (b - a) / a], [None, INT, INT, INT, PCT])
    s.gap()
    s.section("Where the change came from (Year B minus Year A NPR)")
    s.header(["Group", "Customers", "Year A NPR", "Year B NPR", "Change in NPR"])
    st = q1["by_customer_status"]["entries"]
    names = {"retained": "Retained customers", "new_in_b": "New customers", "lost": "Lost customers",
             "cancel_only_no_sale_in_a_or_b": "Cancel-only customers", "no_customer_id": "No customer ID"}
    for k, lab in names.items():
        a, b = num(st[k]["npr_a_gbp"]), num(st[k]["npr_b_gbp"])
        s.row([lab, st[k]["customers"] or "", a, b, num(st[k]["delta_gbp"])], [None, INT, GBP, GBP, GBP])
    s.row(["Total", "", num(q1["npr_a_gbp"]), num(q1["npr_b_gbp"]), num(q1["delta_npr_gbp"])], [None, INT, GBP, GBP, GBP], total=True)
    s.gap()
    s.section("Concentration and repeat purchase")
    s.header(["Metric", "Value", "", "", ""])
    c = R["q2"]["concentration_year_b"]
    s.row(["Top 1% of identified customers: share of identified Year B NPR", c["top_1pct"]["share_of_identified_npr"]], [None, PCT])
    s.row(["Top 10 customers: share of identified Year B NPR", c["top_10_customers"]["share_of_identified_npr"]], [None, PCT])
    s.row([f"Pooled 90-day repeat rate (n = {F['repeat_n']:,}, cohorts {F['repeat_first']} to {F['repeat_last']})", F["repeat_rate"]], [None, PCT])
    s.row(["Cancelled value traceable to an earlier order (whole period)", F["traceable_value_share"]], [None, PCT])
    s.gap()
    s.section("The story")
    A = answers(R, F)
    for k in ("q1", "q1_country", "q2", "q3"):
        s.text(A[k], wrap_cols=5, height=height_for(A[k], 118))
    s.gap()
    s.text(SOURCE_LINE + ". " + DEMO_NOTE + ".", f_sub, wrap_cols=5)


def sheet_monthly(wb, R, F):
    q1, q3 = R["q1"], R["q3"]
    s = Sheet(wb, "Monthly", [13, 8, 9, 16, 16, 16, 12, 3, 12, 16, 16, 14], tab=S.NEUTRAL.lstrip("#"))
    s.title("Monthly NPR, GPS and CPV", "All 25 months. 2011-12 is partial (1-9 Dec) and excluded from every year comparison.")
    hdr = s.header(["Month", "Year", "Partial", "GPS", "CPV", "NPR", "CPV / GPS"])
    first = s.r
    for m, v in q1["monthly"].items():
        s.row([m, v["year"] or "-", "yes" if v["partial"] else "", num(v["gps_gbp"]), num(v["cpv_gbp"]), num(v["npr_gbp"]),
               q3["monthly"][m]["cpv_over_gps"]], [None, None, None, GBP, GBP, GBP, PCT])
    last = s.r - 1
    s.freeze(hdr + 1)
    # Year A vs Year B side by side (feeds the chart)
    ws = s.ws
    r0 = hdr
    for j, c in enumerate(["Month of year", "Year A NPR", "Year B NPR", "B minus A"], 9):
        cell = ws.cell(r0, j, c)
        cell.font, cell.fill = f_head, fill_head
        cell.alignment = Alignment(horizontal="left" if j == 9 else "right", vertical="center", wrap_text=True)
    for i, (lab, a, b) in enumerate(year_pairs(R), 1):
        va, vb = num(q1["monthly"][a]["npr_gbp"]), num(q1["monthly"][b]["npr_gbp"])
        for j, (v, fmt) in enumerate(((lab, None), (va, GBP), (vb, GBP), (vb - va, GBP)), 9):
            cell = ws.cell(r0 + i, j, v)
            cell.font, cell.border = f_body, b_bottom
            if fmt:
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right")
    tot = r0 + 13
    for j, (v, fmt) in enumerate((("Year total", None), (num(F["npr_a"]), GBP), (num(F["npr_b"]), GBP), (num(F["delta"]), GBP)), 9):
        cell = ws.cell(tot, j, v)
        cell.font, cell.fill, cell.border = f_bold, fill_total, b_top
        if fmt:
            cell.number_format = fmt
    ch = LineChart()
    ch.title = f"NPR by month: Year B vs Year A"
    ch.height, ch.width = 8.5, 17
    data = Reference(ws, min_col=10, max_col=11, min_row=r0, max_row=r0 + 12)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=9, min_row=r0 + 1, max_row=r0 + 12))
    for ser, colour in zip(ch.series, (S.YEAR_A.lstrip("#"), INK)):
        ser.graphicalProperties.line.solidFill = colour
        ser.graphicalProperties.line.width = 28000
        ser.smooth = False
        ser.marker.symbol = "circle"
        ser.marker.size = 5
        ser.marker.graphicalProperties.solidFill = colour
        ser.marker.graphicalProperties.line.solidFill = colour
    ch.y_axis.number_format = '£#,##0'
    ch.y_axis.majorGridlines.spPr = None
    ch.y_axis.delete = ch.x_axis.delete = False
    ch.legend.position = "b"
    ws.add_chart(ch, f"I{tot + 3}")


def sheet_growth(wb, R, F):
    q1 = R["q1"]
    s = Sheet(wb, "Growth", [44, 15, 17, 17, 17, 13], tab=S.GROWTH.lstrip("#"))
    s.title("Where did revenue growth come from?", f"NPR Year B minus Year A = {num(q1['delta_npr_gbp']):,.0f} GBP ({F['growth'] * 100:.1f}%). Each decomposition sums to that change exactly.")
    def block(title, note, rows, total_label="Total change in NPR"):
        s.section(title, note)
        s.header(["", "Customers", "Year A NPR", "Year B NPR", "Change", "Change %"])
        for lab, cust, a, b in rows:
            a, b = num(a), num(b)
            s.row([lab, cust, a, b, b - a, (b - a) / a if a > 0 else None], [None, INT, GBP, GBP, GBP, PCT])
        ta, tb = sum(num(x[2]) for x in rows), sum(num(x[3]) for x in rows)
        s.row([total_label, None, ta, tb, tb - ta, (tb - ta) / ta], [None, INT, GBP, GBP, GBP, PCT], total=True)
        s.gap()
    ce = q1["by_country_group"]["entries"]
    block("Country group", "United Kingdom, the five largest non-UK countries by Year B NPR, other.",
          [(k, None, v["npr_a_gbp"], v["npr_b_gbp"]) for k, v in ce.items()])
    st = q1["by_customer_status"]["entries"]
    names = {"retained": "Retained (bought in A and B)", "new_in_b": "New (bought in B only)", "lost": "Lost (bought in A only)",
             "cancel_only_no_sale_in_a_or_b": "Cancel-only (no sale in A or B)", "no_customer_id": "No customer ID"}
    block("Customer status", "A customer's change includes its cancellations. Change % is left blank where the Year A base is not positive.",
          [(names[k], v["customers"], v["npr_a_gbp"], v["npr_b_gbp"]) for k, v in st.items()])
    bp = q1["by_product"]
    s.section("Top 10 products by absolute change",
              f"{bp['products']:,} products in total. These ten account for {bp['top10_share_of_gross_movement'] * 100:.1f}% of the total gross movement "
              f"(sum of absolute changes {num(bp['total_gross_movement_gbp']):,.0f} GBP); their net change is {num(bp['top10_net_delta_gbp']):,.0f} GBP.")
    s.header(["Product", "Stock code", "Year A NPR", "Year B NPR", "Change", ""])
    for p in bp["top10"]:
        a, b = num(p["npr_a_gbp"]), num(p["npr_b_gbp"])
        s.row([p["description"].strip(), p["stock_code"], a, b, b - a], [None, None, GBP, GBP, GBP])
    s.gap()
    br = q1["bridge"]
    s.section("Bridge: identified-customer NPR = customers x orders per customer x NPR per order",
              "Plus NPR without a customer ID. Orders are those of identified customers.")
    s.header(["", "Year A", "Year B", "B / A", "", ""])
    a, b, r = br["year_a"], br["year_b"], br["ratio_b_over_a"]
    s.row(["Identified customers", a["customers"], b["customers"], r["customers"]], [None, INT, INT, "0.000"])
    s.row(["Orders", a["orders"], b["orders"], b["orders"] / a["orders"]], [None, INT, INT, "0.000"])
    s.row(["Orders per customer", a["orders_per_customer"], b["orders_per_customer"], r["orders_per_customer"]], [None, "0.00", "0.00", "0.000"])
    s.row(["NPR per order", num(a["npr_per_order_gbp"]), num(b["npr_per_order_gbp"]), r["npr_per_order"]], [None, GBP2, GBP2, "0.000"])
    s.row(["Identified NPR", num(a["identified_npr_gbp"]), num(b["identified_npr_gbp"]), r["identified_npr"]], [None, GBP, GBP, "0.000"], bold=True)
    s.row(["No-customer-ID NPR", num(a["no_id_npr_gbp"]), num(b["no_id_npr_gbp"]), num(b["no_id_npr_gbp"]) / num(a["no_id_npr_gbp"])], [None, GBP, GBP, "0.000"])
    s.row(["Total NPR", num(a["npr_gbp"]), num(b["npr_gbp"]), num(b["npr_gbp"]) / num(a["npr_gbp"])], [None, GBP, GBP, "0.000"], total=True)
    s.freeze(4)


def sheet_customers(wb, R, F):
    q2 = R["q2"]
    conc = q2["concentration_year_b"]
    s = Sheet(wb, "Customers", [30, 13, 13, 13, 13, 13, 13, 13, 13, 13, 13, 13, 13, 13], tab=S.NO_ID.lstrip("#"))
    s.title("How dependent is revenue on a few customers?", "Year B, identified customers only. Customers with negative NPR are kept.")
    s.section("Concentration", f"{conc['identified_customers']:,} identified customers; {conc['negative_npr_customers']} have negative NPR.")
    s.header(["Group", "Customers", "NPR", "Share of identified NPR", "Share of total NPR"])
    for key, lab in (("top_1pct", "Top 1%"), ("top_10pct", "Top 10%"), ("top_20pct", "Top 20%"), ("top_10_customers", "Top 10 customers")):
        c = conc[key]
        s.row([lab, c["customers"], num(c["npr_gbp"]), c["share_of_identified_npr"], c["share_of_total_npr"]], [None, INT, GBP, PCT, PCT])
    s.row(["All identified customers", conc["identified_customers"], num(conc["identified_npr_gbp"]), 1.0, num(conc["identified_npr_gbp"]) / num(conc["total_npr_gbp"])],
          [None, INT, GBP, PCT, PCT], total=True)
    s.gap()
    rp = q2["repeat_purchase"]
    s.section("Repeat purchase: second order within 90 days of the first",
              "Cohort = month of a customer's first order in the data. Cohorts before 2010-03 are left-censored; later cohorts need a full 90-day window.")
    s.header(["Cohort", "Customers", "Repeaters", "Repeat rate", "In average?"])
    for c, v in rp["cohorts"].items():
        s.row([c, v["customers"], v["repeaters"], v["rate"], "yes" if v["included"] else "no"], [None, INT, INT, PCT])
    p = rp["pooled"]
    s.row(["Pooled (included cohorts)", p["customers"], p["repeaters"], p["rate"], ""], [None, INT, INT, PCT], total=True)
    pre = rp["before_2010_03_shown_separately"]
    s.row(["Cohorts before 2010-03 (shown separately)", pre["customers"], pre["repeaters"], pre["pooled_rate"], "no"], [None, INT, INT, PCT])
    s.gap()
    s.section("Monthly retention matrix", "Share of the cohort with at least one order in month +N. Blank = beyond the last full month (Nov 2011).")
    offs = q2["retention_offsets"]
    hdr = s.header(["Cohort"] + [f"+{o}" for o in offs])
    first = s.r
    for c, row in q2["retention_matrix"].items():
        s.row([c] + row, [None] + [PCT] * len(offs))
    last = s.r - 1
    s.ws.conditional_formatting.add(f"C{first}:{get_column_letter(1 + len(offs))}{last}",
                                    ColorScaleRule(start_type="num", start_value=0, start_color="FFFEFB",
                                                   end_type="num", end_value=0.35, end_color=S.GROWTH.lstrip("#").upper()))
    s.freeze(4)


def sheet_cancellations(wb, R, F):
    q3 = R["q3"]
    s = Sheet(wb, "Cancellations", [40, 14, 14, 14, 14, 14, 14, 16, 14], tab=S.LOSS.lstrip("#"))
    s.title("How much revenue is lost to cancellations?", "CPV = cancelled product value, GPS = gross product sales. Every headline is shown with and without the 10 largest single cancellation lines.")
    s.section("Headlines by period")
    s.header(["Period / variant", "GPS", "CPV", "CPV / GPS", "Cancel lines", "Top-20 products share of CPV", "Top-20 customers share of CPV", "Traceable share of CPV value", ""])
    for per, lab in (("year_a", "Year A"), ("year_b", "Year B"), ("whole_period", "Whole period (incl. Dec 2011)")):
        for var, vl in (("with_top10_lines", "with all lines"), ("without_top10_lines", "without the 10 largest lines")):
            h = q3["headlines"][per][var]
            s.row([f"{lab}, {vl}", num(h["gps_gbp"]), num(h["cpv_gbp"]), h["cpv_over_gps"], h["cancel_lines"], h["top20_products_share_of_cpv"],
                   h["top20_customers_share_of_cpv"], h["traceability"]["traceable"]["share_of_value"]],
                  [None, GBP, GBP, PCT, INT, PCT, PCT, PCT])
    s.gap()
    s.section("Monthly CPV with and without the 10 largest lines", "2011-12 is partial (1-9 Dec).")
    s.header(["Month", "GPS", "CPV", "CPV / GPS", "CPV without top 10", "CPV / GPS without", "Top-10 lines value", "Partial", ""])
    for m, v in q3["monthly"].items():
        w = q3["monthly_without_top10_lines"][m]
        s.row([m, num(v["gps_gbp"]), num(v["cpv_gbp"]), v["cpv_over_gps"], num(w["cpv_gbp"]), w["cpv_over_gps"],
               num(v["cpv_gbp"]) - num(w["cpv_gbp"]), "yes" if v["partial"] else ""], [None, GBP, GBP, PCT, GBP, PCT, GBP])
    s.gap()
    t = q3["top10_cancellation_lines"]
    s.section("The 10 largest single cancellation lines", f"Together {num(t['total_value_gbp']):,.0f} GBP, {t['share_of_whole_period_cpv'] * 100:.1f}% of whole-period CPV.")
    s.header(["Product", "Invoice", "Date", "Customer ID", "Quantity", "Price", "Value", "Traceable", "Country"])
    for ln in t["lines"]:
        s.row([ln["description"].strip(), ln["invoice"], ln["date"][:10], ln["customer_id"], ln["quantity"], num(ln["price_gbp"]), num(ln["value_gbp"]),
               "yes" if ln["traceable"] else "no", ln["country"]], [None, None, None, "0", INT, GBP2, GBP])
    s.gap()
    cr = q3["cancellation_rate"]
    s.section("Top 20 products by cancellation rate (cancelled units / sold units)",
              f"Products with at least {cr['min_sold_units']} sold units and {cr['min_sale_invoices']} sale invoices ({cr['with_top10_lines']['eligible_products']:,} eligible). "
              f"Pooled rate {cr['with_top10_lines']['pooled_rate_eligible_products'] * 100:.1f}%, {cr['without_top10_lines']['pooled_rate_eligible_products'] * 100:.1f}% without the 10 largest lines. "
              "Left: all lines. Right block: without the 10 largest lines.")
    s.header(["Product (all lines)", "Stock code", "Sold units", "Cancelled units", "Rate", "Stock code (without top 10)", "Cancelled units", "Rate", ""])
    a, b = cr["with_top10_lines"]["top20"], cr["without_top10_lines"]["top20"]
    for x, y in zip(a, b):
        s.row([x["description"].strip(), x["stock_code"], x["sold_units"], x["cancelled_units"], x["rate"], y["stock_code"], y["cancelled_units"], y["rate"]],
              [None, None, INT, INT, PCT, None, INT, PCT])


def sheet_ledger(wb, R, F):
    rec = R["rec"]
    s = Sheet(wb, "Ledger", [26, 14, 18, 20, 22], tab=S.NEUTRAL.lstrip("#"))
    s.title("Row ledger: every raw row has exactly one disposition", "Rules are applied in a fixed order and the first match wins (see ANALYSIS_PLAN.md).")
    hdr = s.header(["Disposition", "Rows", "Value (GBP)", "Rows without customer ID", "Value without customer ID (GBP)"])
    for r in ledger_rows(R):
        s.row([r["disposition"], r["rows"], num(r["value_gbp"]), r["no_id_rows"], num(r["no_id_value_gbp"])], [None, INT, GBP2, INT, GBP2])
    tot = ledger_total(R)
    s.row(["Raw total", tot["rows"], num(tot["value_gbp"]), tot["no_id_rows"], num(tot["no_id_value_gbp"])], [None, INT, GBP2, INT, GBP2], total=True)
    s.gap()
    s.section("Identity checks")
    s.header(["Check", "Result", "", "", ""])
    for k, v in rec["identities"].items():
        s.row([k.replace("_", " "), "pass" if v else "FAIL"])
    s.row(["Rows in disposition 'unclassified'", rec["dispositions"]["unclassified"]["rows"]], [None, INT])
    s.gap()
    s.section("Cross-sheet overlap")
    ov = rec["sheets"]["overlap_window"]
    s.text(f"The two source sheets overlap between {ov['from']} and {ov['to']}: {ov['rows_sheet2_in_window']:,} rows of sheet 2 are copies of sheet 1 rows "
           f"and are marked dup_cross_sheet ({ov['dup_cross_sheet_rows']:,} rows).", wrap_cols=5, height=32)
    s.freeze(hdr + 1)


def sheet_checks(wb, R, F):
    ic = R["indep"]
    s = Sheet(wb, "Checks", [58, 16, 30], tab=S.NEUTRAL.lstrip("#"))
    s.title("Independent check", "A second implementation, written from the raw CSV and the plan only, re-derived the ledger and the headline numbers.")
    s.header(["Check", "Value", ""])
    s.row(["Raw rows compared", ic["rows_compared"]], [None, INT])
    s.row(["Rows present in only one implementation", ic["rows_only_in_one"]], [None, INT])
    s.row(["Row disposition mismatches", ic["row_disposition_mismatches"]], [None, INT])
    s.row(["Headline numbers compared", ic["numbers_compared"]], [None, INT])
    s.row(["Headline number mismatches", len(ic["number_mismatches"])], [None, INT])
    s.row(["Result", "all agree" if (ic["rows_only_in_one"] == ic["row_disposition_mismatches"] == len(ic["number_mismatches"]) == 0) else "MISMATCH"], bold=True)
    s.gap()
    s.section("Decomposition identities (results/q1.json)")
    s.header(["Identity", "Result", ""])
    for k, v in R["q1"]["identities"].items():
        s.row([k.replace("_", " "), "pass" if v else "FAIL"])
    s.gap()
    s.text("Differences between the two implementations are resolved and documented (ANALYSIS_PLAN.md, Amendments), not averaged.", f_sub, wrap_cols=3, height=30)


def build(out: Path | None = None) -> Path:
    R = load_all()
    F = facts(R)
    wb = Workbook()
    wb.remove(wb.active)
    for fn in (sheet_readme, sheet_summary, sheet_monthly, sheet_growth, sheet_customers, sheet_cancellations, sheet_ledger, sheet_checks):
        fn(wb, R, F)
    wb.properties.title = "Retail sales report (demo)"
    wb.properties.creator = "Elmar Guliyev"
    out = Path(out) if out else REPORTS / "retail-sales-report.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("wrote", out)
    return out


if __name__ == "__main__":
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else None)
