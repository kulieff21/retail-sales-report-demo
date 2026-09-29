"""Dump each sheet of the raw UCI workbook to CSV with no transformation.

Cell values are written as they come out of openpyxl: numbers via repr-free str(),
datetimes as ISO 8601, empty cells as empty strings. The CSVs exist so that tools
without an xlsx reader (e.g. the independent reproduction) read the same raw rows.

Usage: uv run python tools/export_raw_csv.py  (reads $RETAIL_DATA_DIR/online_retail_II.xlsx)
"""

import csv
import datetime as dt
import hashlib
import json
import os
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_XLSX_SHA256 = "bcbe73b35f5b7babf197fb0cb983a11f5d9ff929078d4aa53d171b1f2df2e980"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    return str(value)


def main() -> None:
    data_dir = Path(os.environ["RETAIL_DATA_DIR"])
    xlsx = data_dir / "online_retail_II.xlsx"
    digest = sha256(xlsx)
    if digest != EXPECTED_XLSX_SHA256:
        raise SystemExit(f"unexpected workbook hash {digest}")

    wb = load_workbook(xlsx, read_only=True)
    manifest = {"source": xlsx.name, "sha256": digest, "sheets": []}
    for i, ws in enumerate(wb.worksheets, start=1):
        out = data_dir / f"raw_sheet{i}.csv"
        rows = 0
        with out.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for j, row in enumerate(ws.iter_rows(values_only=True)):
                writer.writerow([cell_to_text(v) for v in row])
                rows += j > 0
        manifest["sheets"].append(
            {"index": i, "title": ws.title, "csv": out.name, "data_rows": rows, "csv_sha256": sha256(out)}
        )
        print(ws.title, rows)
    (Path(__file__).resolve().parents[1] / "results" / "raw_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
