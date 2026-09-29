"""Stage 3a tests: the workbook opens and its key cells equal the JSON values; figures, slide and notebook are produced."""

import json
import os
import sys
from decimal import Decimal
from pathlib import Path

import matplotlib.image as mpimg
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

import build_figures  # noqa: E402
import build_slide  # noqa: E402
import build_workbook  # noqa: E402
from retail_report.report_data import facts, load_all, pct  # noqa: E402

DATA = os.environ.get("RETAIL_DATA_DIR")
has_ledger = pytest.mark.skipif(not DATA or not (Path(DATA) / "ledger.parquet").exists(), reason="ledger.parquet missing")


def num(x):
    return float(Decimal(str(x)))


def find(ws, label, col=1, start=1):
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, col).value == label:
            return r
    raise AssertionError(f"{label!r} not found in {ws.title}")


@pytest.fixture(scope="module")
def R():
    return load_all()


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    path = build_workbook.build(tmp_path_factory.mktemp("xlsx") / "report.xlsx")
    return load_workbook(path)


def test_workbook_sheets_and_layout(wb):
    assert wb.sheetnames == ["README", "Summary", "Monthly", "Growth", "Customers", "Cancellations", "Ledger", "Checks"]
    for name in wb.sheetnames:
        assert wb[name].sheet_view.showGridLines is False
    for name in ("Monthly", "Growth", "Customers", "Ledger"):
        assert wb[name].freeze_panes is not None
    assert len(wb["Monthly"]._charts) == 1
    assert len(wb["Customers"].conditional_formatting) == 1
    readme = " ".join(str(c.value) for row in wb["README"].iter_rows() for c in row if c.value)
    assert "CC BY 4.0" in readme and "Demo analysis on public data" in readme and "10.24432/C5CG6D" in readme


def test_workbook_summary_matches_json(wb, R):
    ws = wb["Summary"]
    r = find(ws, "Net product revenue (NPR)")
    assert ws.cell(r, 2).value == pytest.approx(num(R["q1"]["npr_a_gbp"]))
    assert ws.cell(r, 3).value == pytest.approx(num(R["q1"]["npr_b_gbp"]))
    assert ws.cell(r, 4).value == pytest.approx(num(R["q1"]["delta_npr_gbp"]))
    assert ws.cell(r, 2).number_format.startswith("£")
    r = find(ws, "CPV / GPS")
    assert ws.cell(r, 3).value == R["q3"]["headlines"]["year_b"]["with_top10_lines"]["cpv_over_gps"]
    r = find(ws, "CPV / GPS without the 10 largest lines")
    assert ws.cell(r, 3).value == R["q3"]["headlines"]["year_b"]["without_top10_lines"]["cpv_over_gps"]
    st = R["q1"]["by_customer_status"]["entries"]["retained"]
    r = find(ws, "Retained customers")
    assert ws.cell(r, 5).value == pytest.approx(num(st["delta_gbp"]))


def test_workbook_monthly_growth_customers(wb, R):
    ws = wb["Monthly"]
    r = find(ws, "2011-01")
    assert ws.cell(r, 4).value == pytest.approx(num(R["q1"]["monthly"]["2011-01"]["gps_gbp"]))
    assert ws.cell(r, 6).value == pytest.approx(num(R["q1"]["monthly"]["2011-01"]["npr_gbp"]))
    assert sum(1 for row in ws.iter_rows(min_row=5, max_row=29, min_col=1, max_col=1) if row[0].value) == 25
    ws = wb["Growth"]
    r = find(ws, "Australia")
    assert ws.cell(r, 5).value == pytest.approx(num(R["q1"]["by_country_group"]["entries"]["Australia"]["delta_gbp"]))
    ws = wb["Customers"]
    r = find(ws, "Top 1%")
    assert ws.cell(r, 4).value == R["q2"]["concentration_year_b"]["top_1pct"]["share_of_identified_npr"]
    r = find(ws, "Pooled (included cohorts)")
    assert ws.cell(r, 4).value == R["q2"]["repeat_purchase"]["pooled"]["rate"]
    r = find(ws, "2010-03", start=find(ws, "Cohort", start=find(ws, "Cohort") + 1))
    assert ws.cell(r, 3).value == R["q2"]["retention_matrix"]["2010-03"][1]
    assert ws.cell(r, 15).value is not None or R["q2"]["retention_matrix"]["2010-03"][12] is not None


def test_workbook_cancellations_ledger_checks(wb, R):
    ws = wb["Cancellations"]
    r = find(ws, "Year B, without the 10 largest lines")
    assert ws.cell(r, 3).value == pytest.approx(num(R["q3"]["headlines"]["year_b"]["without_top10_lines"]["cpv_gbp"]))
    r = find(ws, "2011-01", start=find(ws, "Month"))
    assert ws.cell(r, 5).value == pytest.approx(num(R["q3"]["monthly_without_top10_lines"]["2011-01"]["cpv_gbp"]))
    ws = wb["Ledger"]
    r = find(ws, "Raw total")
    assert ws.cell(r, 2).value == R["rec"]["raw"]["rows"]
    assert ws.cell(r, 3).value == pytest.approx(num(R["rec"]["raw"]["value_gbp"]))
    r = find(ws, "sale_product")
    assert ws.cell(r, 2).value == R["rec"]["dispositions"]["sale_product"]["rows"]
    ws = wb["Checks"]
    assert ws.cell(find(ws, "Headline numbers compared"), 2).value == R["indep"]["numbers_compared"]
    assert ws.cell(find(ws, "Result"), 2).value == "all agree"


def test_slide_png_and_pdf(tmp_path):
    png, pdf = build_slide.main(tmp_path)
    img = mpimg.imread(png)
    assert img.shape[:2] == (1080, 1920)
    assert pdf.read_bytes()[:4] == b"%PDF"


@has_ledger
def test_figures_are_produced(tmp_path):
    files = build_figures.main(tmp_path)
    names = {f.name for f in files}
    assert names == {"01_monthly_npr.png", "02_npr_bridge.png", "03_country_delta.png", "04_concentration.png",
                     "05_cohort_retention.png", "06_cancellations.png"}
    assert mpimg.imread(tmp_path / "01_monthly_npr.png").shape[:2] == (900, 1600)   # 8 x 4.5 in at 2x (200 dpi)
    assert all(f.stat().st_size > 20_000 for f in files)


def test_notebook_executed_and_numbers_come_from_code():
    path = ROOT / "notebooks" / "report.ipynb"
    if not path.exists():
        pytest.skip("run tools/build_reports.py first")
    nb = json.loads(path.read_text(encoding="utf-8"))
    for c in nb["cells"]:
        if c["cell_type"] == "markdown":
            src = "".join(c["source"])
            assert "£" not in src and "%" not in src, "numbers must be rendered from JSON, not typed into markdown"
        else:
            assert c["execution_count"] is not None
            assert all(o["output_type"] != "error" for o in c["outputs"])
    text = json.dumps(nb, ensure_ascii=False)
    F = facts(load_all())
    assert pct(F["growth"]) in text and pct(F["retained_pct"], 1, True) in text    # rendered from JSON by code cells
